"""
Tests for how decoding results reach the user: the Inspector summary, the
exports, the chain-of-custody log, the CLI and the signed report.

The rule these enforce is that a reading never travels without its
confidence. A PROBABLE result copied out of a spreadsheet cell, a CSV, an
HTML report or a signed report must still say PROBABLE.
"""

import csv
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import fitz
import pikepdf
from openpyxl import load_workbook

from src import cid_report
from src.exporter import (
    cid_columns_for_path, export_to_csv, export_to_excel, export_to_html,
    export_to_json,
)
from src.signed_report import build_findings_report
from src.touchup_extract import extract_and_decode, extract_touchup_text


def build_touchup_pdf(strip_tounicode=True):
    """
    A PDF with a genuine TouchUp_TextEdit marked-content region.

    Uses the bundled reference font because PyMuPDF's subsetter drops its
    TrueType ``post`` names. With the ToUnicode CMap also removed, nothing is
    left but the shapes, so decoding goes through Tier 2 and yields a
    PROBABLE reading - which is what the labelling rules need to be tested
    against. The base-14 fonts keep a CFF charset and would resolve at Tier 1
    as CERTAIN, leaving those rules unexercised.
    """
    from src.cid_shapes import bundled_reference_font

    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="F1", fontfile=str(bundled_reference_font()))
    page.insert_text((50, 100), "Untouched original line",
                     fontname="F1", fontsize=12)
    page.insert_text((50, 140), "Beloeb 12.450 kr", fontname="F1", fontsize=12)
    doc.subset_fonts()
    data = doc.tobytes()
    doc.close()

    pdf = pikepdf.open(io.BytesIO(data))
    try:
        page = pdf.pages[0]
        new_ops = []
        seen = 0
        for operands, operator in pikepdf.parse_content_stream(page):
            if str(operator) in {"Tj", "TJ"}:
                seen += 1
                if seen == 2:
                    new_ops.append(([pikepdf.Name("/Span"), pikepdf.Name("/TU")],
                                    pikepdf.Operator("BDC")))
                    new_ops.append((operands, operator))
                    new_ops.append(([], pikepdf.Operator("EMC")))
                    continue
            new_ops.append((operands, operator))
        page.Contents = pdf.make_stream(pikepdf.unparse_content_stream(new_ops))

        if "/Resources" not in page:
            page.Resources = pikepdf.Dictionary()
        page.Resources[pikepdf.Name("/Properties")] = pikepdf.Dictionary(
            TU=pdf.make_indirect(pikepdf.Dictionary(
                Type=pikepdf.Name("/TouchUp_TextEdit"))))

        if strip_tounicode:
            for obj in pdf.objects:
                try:
                    if isinstance(obj, pikepdf.Dictionary) and obj.get("/Type") == "/Font":
                        if "/ToUnicode" in obj:
                            del obj["/ToUnicode"]
                except Exception:
                    pass

        out = io.BytesIO()
        pdf.save(out)
    finally:
        pdf.close()
    return out.getvalue()


def decoded_for(strip_tounicode=True):
    data = build_touchup_pdf(strip_tounicode)
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        page_text, runs, custody = extract_and_decode(doc)
    finally:
        doc.close()
    return page_text, runs, custody


def scan_data_with(runs, custody, path="C:/cases/invoice.pdf"):
    return {path: {
        "path": path,
        "indicator_keys": {"TouchUp_TextEdit": {
            "found_text": {1: ["garbled"]},
            "decoded_runs": runs,
            "decode_custody": custody,
        }},
    }}


COLUMN_KEYS = ["col_id", "col_name", "col_altered", "col_revisions", "col_path",
               "col_created", "col_modified", "col_md5", "col_exif",
               "col_indicators", "col_note"]


# --------------------------------------------------------------------------
# The extraction path itself
# --------------------------------------------------------------------------

class TestTouchUpExtraction(unittest.TestCase):
    def test_masking_isolates_the_touchup_region(self):
        """Only the marked text survives; the rest must be gone."""
        page_text, _runs, _custody = decoded_for(strip_tounicode=False)
        joined = " ".join(" ".join(v) for v in page_text.values())
        self.assertIn("Beloeb", joined)
        self.assertNotIn("Untouched", joined)

    def test_masking_leaves_no_substitution_artefacts(self):
        """
        Masked two-byte text must not come back as daggers.

        Overwriting a two-byte encoding with 0x20 bytes yields code 0x2020,
        which renders as U+2020 DAGGER.
        """
        page_text, _runs, _custody = decoded_for(strip_tounicode=False)
        joined = " ".join(" ".join(v) for v in page_text.values())
        self.assertNotIn("\u2020", joined)

    def test_runs_are_captured_with_their_font(self):
        data = build_touchup_pdf()
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            _text, runs, _pdf_bytes = extract_touchup_text(doc, capture_runs=True)
        finally:
            doc.close()
        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0]["page"], 1)
        self.assertTrue(runs[0]["font_resource"])
        self.assertTrue(runs[0]["encoded"])

    def test_decoding_recovers_text_without_a_cmap(self):
        _text, runs, custody = decoded_for(strip_tounicode=True)
        self.assertEqual(len(runs), 1)
        self.assertIsNotNone(runs[0]["text"])
        self.assertIn(runs[0]["confidence"], ("CERTAIN", "PROBABLE"))
        self.assertTrue(custody)

    def test_scan_paths_share_one_implementation(self):
        """
        The worker and in-process paths must not drift.

        They were separate copies; a fix to one did not reach the other, so
        the same file could give different TouchUp text depending on which
        path scanned it.
        """
        from src import scan_worker, touchup_extract
        from src.data_processing import DataProcessingMixin

        data = build_touchup_pdf()

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            worker_text = scan_worker._extract_touchup_text(doc)
        finally:
            doc.close()

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            mixin_text = DataProcessingMixin._extract_touchup_text(None, doc)
        finally:
            doc.close()

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            shared_text = touchup_extract.extract_touchup_text(doc)
        finally:
            doc.close()

        self.assertEqual(worker_text, shared_text)
        self.assertEqual(mixin_text, shared_text)


# --------------------------------------------------------------------------
# Presentation helpers
# --------------------------------------------------------------------------

class TestPresentation(unittest.TestCase):
    def setUp(self):
        _text, self.runs, self.custody = decoded_for()

    def test_certain_text_is_unadorned(self):
        record = {"text": "Beloeb", "confidence": "CERTAIN"}
        self.assertEqual(cid_report.labelled_text(record), "Beloeb")

    def test_lesser_confidence_is_always_labelled(self):
        for level in ("PROBABLE", "SPECULATIVE"):
            record = {"text": "Beloeb", "confidence": level}
            self.assertTrue(cid_report.labelled_text(record).startswith(f"[{level}]"))

    def test_missing_text_is_not_silently_empty(self):
        record = {"text": None, "confidence": "SPECULATIVE", "error": "boom"}
        rendered = cid_report.labelled_text(record)
        self.assertIn("SPECULATIVE", rendered)
        self.assertIn("boom", rendered)

    def test_weakest_confidence_wins(self):
        records = [{"confidence": "CERTAIN"}, {"confidence": "PROBABLE"}]
        self.assertEqual(cid_report.weakest_confidence(records), "PROBABLE")
        records.append({"confidence": "SPECULATIVE"})
        self.assertEqual(cid_report.weakest_confidence(records), "SPECULATIVE")

    def test_unknown_confidence_is_treated_as_weakest(self):
        """An unrecognised level must never be read as confirmed."""
        self.assertEqual(
            cid_report.weakest_confidence([{"confidence": "banana"}]),
            "SPECULATIVE")

    def test_export_rows_carry_verifiable_facts(self):
        row = cid_report.export_rows(self.runs)[0]
        self.assertTrue(row["encoded_hex"])
        self.assertEqual(len(row["font_sha256"]), 64)
        self.assertIn(row["confidence"], ("CERTAIN", "PROBABLE"))
        self.assertTrue(row["tiers"])

    def test_custody_details_include_hashes(self):
        details = cid_report.custody_details(self.custody, self.runs)
        self.assertIn("tiers", details)
        self.assertIn("weakest_confidence", details)
        self.assertTrue(details["document_font_sha256"])


# --------------------------------------------------------------------------
# Exports
# --------------------------------------------------------------------------

class TestExports(unittest.TestCase):
    def setUp(self):
        _text, self.runs, self.custody = decoded_for()
        self.path = "C:/cases/invoice.pdf"
        self.scan_data = scan_data_with(self.runs, self.custody, self.path)
        self.report_data = [[1, "invoice.pdf", "YES", 0, self.path,
                             "", "", "", "", "", ""]]
        self.confidence = self.runs[0]["confidence"]
        self.text = self.runs[0]["text"]

    def test_columns_for_a_file_without_decoding_are_blank(self):
        self.assertEqual(cid_columns_for_path({}, "nope.pdf"), ["", "", ""])

    def test_excel_has_summary_columns_and_a_detail_sheet(self):
        with TemporaryDirectory() as tmp:
            out = Path(tmp) / "r.xlsx"
            export_to_excel(out, self.report_data, self.scan_data, {}, {},
                            list(COLUMN_KEYS))
            wb = load_workbook(out)
            self.assertIn("Text Decoding", wb.sheetnames)

            main = wb["PDFRecon Results"]
            headers = [c.value for c in main[1]]
            self.assertEqual(headers[-3:],
                             ["Decoded Text", "Decoding Method", "Decoding Confidence"])
            tail = [c.value for c in main[2]][-3:]
            self.assertIn(self.confidence, tail[2])
            if self.confidence != "CERTAIN":
                self.assertIn(f"[{self.confidence}]", tail[0])

            detail = wb["Text Decoding"]
            self.assertEqual(detail.max_row, 2)

    def test_csv_labels_the_reading(self):
        with TemporaryDirectory() as tmp:
            out = Path(tmp) / "r.csv"
            export_to_csv(out, self.report_data, self.scan_data, {}, {},
                          list(COLUMN_KEYS))
            with open(out, encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
        self.assertEqual(rows[0][-1], "Decoding Confidence")
        self.assertEqual(rows[1][-1], self.confidence)
        if self.confidence != "CERTAIN":
            self.assertIn(f"[{self.confidence}]", rows[1][-3])

    def test_html_marks_confidence_distinctly(self):
        with TemporaryDirectory() as tmp:
            out = Path(tmp) / "r.html"
            export_to_html(out, self.report_data, {}, self.scan_data,
                           list(COLUMN_KEYS))
            html = out.read_text(encoding="utf-8")
        self.assertIn("Decoded Text", html)
        self.assertIn(f"conf-{self.confidence.lower()}", html)
        self.assertIn(".conf-speculative", html)  # style is always defined

    def test_json_surfaces_decoding_at_top_level(self):
        with TemporaryDirectory() as tmp:
            out = Path(tmp) / "r.json"
            export_to_json(out, self.scan_data, {}, {})
            payload = json.loads(out.read_text(encoding="utf-8"))
        record = payload["scan_results"][0]
        self.assertIn("text_decoding", record)
        self.assertEqual(record["text_decoding"]["summary"]["runs"], 1)

    def test_no_export_leaks_an_unlabelled_reading(self):
        """The decoded text must never appear without its confidence."""
        if self.confidence == "CERTAIN":
            self.skipTest("only meaningful for inferred readings")
        with TemporaryDirectory() as tmp:
            csv_path = Path(tmp) / "r.csv"
            html_path = Path(tmp) / "r.html"
            export_to_csv(csv_path, self.report_data, self.scan_data, {}, {},
                          list(COLUMN_KEYS))
            export_to_html(html_path, self.report_data, {}, self.scan_data,
                           list(COLUMN_KEYS))
            for blob in (csv_path.read_text(encoding="utf-8-sig"),
                         html_path.read_text(encoding="utf-8")):
                index = blob.find(self.text)
                self.assertNotEqual(index, -1)
                # The label must immediately precede every occurrence.
                while index != -1:
                    self.assertIn(f"[{self.confidence}]", blob[max(0, index - 40):index])
                    index = blob.find(self.text, index + 1)


# --------------------------------------------------------------------------
# Signed report and custody
# --------------------------------------------------------------------------

class TestSignedReportAndCustody(unittest.TestCase):
    def setUp(self):
        _text, self.runs, self.custody = decoded_for()
        self.path = "C:/cases/invoice.pdf"
        self.scan_data = scan_data_with(self.runs, self.custody, self.path)

    def test_signed_report_carries_decoding_not_just_present(self):
        """
        Other nested indicators collapse to "(present)". Decoding must not:
        a reading in a signed report has to arrive with its evidence.
        """
        report = build_findings_report(self.scan_data, {}, {}, {})
        finding = report["findings"][0]
        self.assertEqual(finding["indicators"]["TouchUp_TextEdit"], "(present)")
        self.assertIn("text_decoding", finding)
        self.assertEqual(finding["text_decoding"]["runs"][0]["text"],
                         self.runs[0]["text"])
        self.assertIn("note", finding["text_decoding"])

    def test_signed_report_omits_decoding_when_there_is_none(self):
        report = build_findings_report(
            {"a.pdf": {"path": "a.pdf", "indicator_keys": {}}}, {}, {}, {})
        self.assertNotIn("text_decoding", report["findings"][0])

    def test_custody_entry_is_written_and_chains(self):
        from src.chain_of_custody import (
            ACTION_TEXT_DECODED, log_text_decoding, read_and_verify_custody_log,
        )
        with TemporaryDirectory() as tmp:
            log = Path(tmp) / "custody.log"
            details = cid_report.custody_details(self.custody, self.runs)
            log_text_decoding(log, Path(self.path), "a" * 64, details)
            entries, valid, bad, _msg = read_and_verify_custody_log(log)
        self.assertTrue(valid, f"chain broken at {bad}")
        self.assertEqual(entries[0]["action"], ACTION_TEXT_DECODED)
        self.assertIn("tiers", entries[0]["details"])
        self.assertEqual(entries[0]["details"]["weakest_confidence"],
                         self.runs[0]["confidence"])


if __name__ == "__main__":
    unittest.main()
